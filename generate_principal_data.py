import os
import json
import openpyxl

XLSX_PATH = r"D:\GSQAC AND SAT\MIS SITE\CTS DATA\PRINCIPAL MOBILE NUMBER.xlsx"
DATA_JSON_PATH = r"D:\GSQAC AND SAT\MIS SITE\data\dashboard_data.json"
DATA_JS_PATH = r"D:\GSQAC AND SAT\MIS SITE\data\dashboard_data.js"
PRINCIPAL_JSON_PATH = r"D:\GSQAC AND SAT\MIS SITE\data\principal_data.json"
PRINCIPAL_JS_PATH = r"D:\GSQAC AND SAT\MIS SITE\data\principal_data.js"

print("Reading principal workbook:", XLSX_PATH)
wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
sheet = wb.active

principals_by_udise = {}
principal_records = []

for r in range(4, sheet.max_row + 1):
    sl_no = sheet.cell(r, 1).value
    dist = sheet.cell(r, 2).value
    block = sheet.cell(r, 3).value
    udise = sheet.cell(r, 4).value
    school = sheet.cell(r, 5).value
    status = sheet.cell(r, 6).value
    user_name = sheet.cell(r, 7).value
    desig = sheet.cell(r, 8).value
    mobile = sheet.cell(r, 9).value
    email = sheet.cell(r, 10).value

    if udise and str(udise).strip():
        u_str = str(udise).strip()
        item = {
            "sl_no": sl_no,
            "udise_code": u_str,
            "school_id": u_str,
            "school_name": str(school).strip() if school else "",
            "status": str(status).strip() if status else "0-Operational",
            "principal_name": str(user_name).strip() if user_name else "",
            "designation": str(desig).strip() if desig else "PRINCIPAL",
            "mobile": str(mobile).strip() if mobile else "",
            "email": str(email).strip() if email else "",
            "district": str(dist).strip() if dist else "MAHESANA",
            "block": str(block).strip() if block else "KADI",
            "cluster": "",
            "management": "",
            "category": ""
        }
        principal_records.append(item)
        principals_by_udise[u_str] = item

print(f"Loaded {len(principal_records)} principal records.")

# Load existing dashboard_data.json to merge cluster & management
dash_data = {}
if os.path.exists(DATA_JSON_PATH):
    with open(DATA_JSON_PATH, "r", encoding="utf-8") as f:
        dash_data = json.load(f)

schools_map = {}
if "school_records" in dash_data:
    for s in dash_data["school_records"]:
        sid = str(s.get("school_id") or s.get("dise_code") or "").strip()
        if sid:
            schools_map[sid] = s

# Enrich principal records with cluster, management, etc.
for p in principal_records:
    sid = p["udise_code"]
    if sid in schools_map:
        sc = schools_map[sid]
        p["cluster"] = sc.get("cluster_name") or sc.get("cluster") or ""
        p["management"] = sc.get("management") or ""
        p["category"] = sc.get("category") or ""
        p["total_students"] = sc.get("total") or 0
        # Also enrich school_records in dash_data!
        sc["principal_name"] = p["principal_name"]
        sc["principal_mobile"] = p["mobile"]
        sc["principal_email"] = p["email"]
        sc["principal_designation"] = p["designation"]

dash_data["principal_records"] = principal_records
dash_data["principals_by_udise"] = principals_by_udise

# Write back dashboard_data.json and dashboard_data.js
with open(DATA_JSON_PATH, "w", encoding="utf-8") as f:
    json.dump(dash_data, f, ensure_ascii=False, indent=2)

with open(DATA_JS_PATH, "w", encoding="utf-8") as f:
    f.write("var globalData = ")
    json.dump(dash_data, f, ensure_ascii=False)
    f.write(";\n")

# Write standalone principal_data.json & principal_data.js
principal_output = {
    "total_principals": len(principal_records),
    "records": principal_records,
    "by_udise": principals_by_udise
}

with open(PRINCIPAL_JSON_PATH, "w", encoding="utf-8") as f:
    json.dump(principal_output, f, ensure_ascii=False, indent=2)

with open(PRINCIPAL_JS_PATH, "w", encoding="utf-8") as f:
    f.write("// Principal Contact Data for Fast Client Load\n")
    f.write("window.principalData = ")
    json.dump(principal_output, f, ensure_ascii=False)
    f.write(";\n")

print("Updated:", DATA_JSON_PATH)
print("Updated:", DATA_JS_PATH)
print("Generated:", PRINCIPAL_JSON_PATH)
print("Generated:", PRINCIPAL_JS_PATH)
